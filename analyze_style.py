#!/usr/bin/env python3
"""
Analyze Portle WEB - Scarico Dati.pptx to extract style and layout information
"""

from pptx import Presentation
from pptx.util import Inches, Pt
import os

# Open the reference presentation
ref_pptx = r'c:\Users\roger.voyat\OneDrive - Accenture\Documents\Claude\workshop\meridian-workshop\Portle WEB - Scarico Dati.pptx'
prs_ref = Presentation(ref_pptx)

print("=" * 80)
print("PORTLE WEB - SCARICO DATI.PPTX STYLE ANALYSIS")
print("=" * 80)

# Analyze slide dimensions
print(f"\nSlide Dimensions: {prs_ref.slide_width} x {prs_ref.slide_height}")
print(f"Width: {prs_ref.slide_width.inches} inches, Height: {prs_ref.slide_height.inches} inches")

# Analyze first few slides
for slide_idx in range(min(5, len(prs_ref.slides))):
    slide = prs_ref.slides[slide_idx]
    print(f"\n--- SLIDE {slide_idx + 1} ---")
    print(f"Slide layout: {slide.slide_layout.name}")
    print(f"Number of shapes: {len(slide.shapes)}")
    
    # Analyze background
    background = slide.background
    fill = background.fill
    print(f"Background fill type: {fill.type}")
    if hasattr(fill, 'fore_color'):
        try:
            if hasattr(fill.fore_color, 'rgb'):
                print(f"Background RGB: {fill.fore_color.rgb}")
        except:
            print("Background color: Unable to extract")
    
    # Analyze shapes (text boxes, rectangles, etc.)
    for shape_idx, shape in enumerate(slide.shapes):
        print(f"\n  Shape {shape_idx}: {shape.name}")
        print(f"    Type: {shape.shape_type}")
        print(f"    Position: ({shape.left.inches:.2f}, {shape.top.inches:.2f})")
        print(f"    Size: {shape.width.inches:.2f} x {shape.height.inches:.2f}")
        
        # Check fill
        if hasattr(shape, 'fill'):
            try:
                if shape.fill.type:
                    print(f"    Fill type: {shape.fill.type}")
                    if hasattr(shape.fill, 'fore_color') and hasattr(shape.fill.fore_color, 'rgb'):
                        print(f"    Fill RGB: {shape.fill.fore_color.rgb}")
            except:
                pass
        
        # Check text
        if hasattr(shape, 'text_frame'):
            text_frame = shape.text_frame
            if text_frame.text:
                print(f"    Text: {text_frame.text[:50]}...")
                for para_idx, para in enumerate(text_frame.paragraphs):
                    if para.text:
                        if para.runs:
                            font = para.runs[0].font
                            print(f"      Paragraph {para_idx}:")
                            print(f"        Font size: {font.size}")
                            print(f"        Font name: {font.name}")
                            print(f"        Bold: {font.bold}")
                            if hasattr(font, 'color') and hasattr(font.color, 'rgb'):
                                try:
                                    print(f"        Color RGB: {font.color.rgb}")
                                except:
                                    pass

print("\n" + "=" * 80)
print(f"Total slides: {len(prs_ref.slides)}")
print("=" * 80)
